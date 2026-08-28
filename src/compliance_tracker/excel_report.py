"""Generates the centralized Excel compliance tracker.

Two sheets: **Tracker** is the primary dashboard -- one row per asset, one
column per configured rule, so a reviewer can scan a single matrix instead of
cross-referencing a summary against a separate issues list. **Audit Log**
keeps the full one-row-per-violation trail for traceability.

Every column beyond the fixed computed ones (Next Deadline, Last Reminder
Sent, Reminder Count, Status) comes from config.excel.info_columns or from
each rule's own label -- this module never hardcodes a field name, so
pointing the whole pipeline at a new domain only requires a new config.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from compliance_tracker.config_schema import AppConfig, ColumnConfig
from compliance_tracker.reminder_log import ReminderSummary
from compliance_tracker.rules import resolve_value
from compliance_tracker.validator import COMPLIANT, FLAGGED, AssetResult

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PASS_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="FCA5A5", end_color="FCA5A5", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")

AUDIT_LOG_COLUMNS = [
    ColumnConfig(field="asset_id", label="Asset ID"),
    ColumnConfig(field="rule_id", label="Rule ID"),
    ColumnConfig(field="severity", label="Severity"),
    ColumnConfig(field="field", label="Field"),
    ColumnConfig(field="value", label="Value"),
    ColumnConfig(field="message", label="Message"),
]

NOTES_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

NEXT_DEADLINE_LABEL = "Next Deadline"
LAST_REMINDER_LABEL = "Last Reminder Sent"
REMINDER_COUNT_LABEL = "Reminder Count"
STATUS_LABEL = "Status"


def _read_existing_notes(path: Path, config: AppConfig) -> dict[str, dict[str, str]]:
    """Read back {asset_id: {notes_label: value}} from a previously-generated
    Tracker sheet, so hand-typed notes survive regeneration. Any failure to
    read (missing sheet, unexpected layout, corrupt file) is treated as "no
    prior notes" rather than raised -- this is a best-effort preservation,
    not a requirement for the run to succeed."""
    notes_labels = [c.label for c in config.excel.notes_columns]
    if not notes_labels or not path.exists():
        return {}

    try:
        wb = load_workbook(path)
        if "Tracker" not in wb.sheetnames:
            return {}
        ws = wb["Tracker"]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = list(header_row)
        # The id column is always info_columns[0], enforced at config-load
        # time to equal source.id_field, so it's always Tracker column A.
        id_col_idx = 0
        label_to_idx = {label: headers.index(label) for label in notes_labels if label in headers}
        if not label_to_idx:
            return {}

        preserved: dict[str, dict[str, str]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if id_col_idx >= len(row) or row[id_col_idx] is None:
                continue
            asset_id = str(row[id_col_idx])
            preserved[asset_id] = {
                label: row[idx] for label, idx in label_to_idx.items() if idx < len(row)
            }
        return preserved
    except Exception:
        return {}


def _write_header(ws: Worksheet, labels: list[str]) -> None:
    for col_idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = "A2"


def _autosize_columns(ws: Worksheet, ncols: int) -> None:
    for col_idx in range(1, ncols + 1):
        max_len = len(str(ws.cell(row=1, column=col_idx).value or ""))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2):
            value = row[0].value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 60)


def next_deadline(config: AppConfig, result: AssetResult) -> str:
    """The earliest date among an asset's not_expired-type rule fields, or ""
    if it has none. Public (not the module's other _-prefixed helpers)
    because email_drafter.py reuses it for the reminder email's deadline
    line -- one source of truth for "what's this asset's next deadline"."""
    dates = []
    for rule in config.rules:
        if rule.type != "not_expired":
            continue
        raw = resolve_value(rule, result.record)
        try:
            dates.append(date.fromisoformat(raw.strip()))
        except (ValueError, AttributeError):
            continue
    return min(dates).isoformat() if dates else ""


def tracker_headers(config: AppConfig) -> list[str]:
    return (
        [c.label for c in config.excel.info_columns]
        + [r.label for r in config.rules]
        + [NEXT_DEADLINE_LABEL, LAST_REMINDER_LABEL, REMINDER_COUNT_LABEL, STATUS_LABEL]
        + [c.label for c in config.excel.notes_columns]
    )


def build_tracker_table(
    config: AppConfig,
    results: list[AssetResult],
    reminder_summary: dict[str, ReminderSummary],
    existing_notes: dict[str, dict[str, str]],
) -> tuple[list[str], list[AssetResult], list[list[str]]]:
    """Presentation-agnostic Tracker content: headers, the same assets in
    display order, and each asset's row values. Used by both the Excel
    writer (which layers styling/fills on top) and the Google Sheets sync
    (which writes the values as-is) -- this is the single place that decides
    what goes in each Tracker cell."""
    info_columns = config.excel.info_columns
    rules = config.rules
    notes_columns = config.excel.notes_columns

    headers = tracker_headers(config)
    ordered = sorted(results, key=lambda r: (r.compliance_status != FLAGGED, r.asset_id))

    rows: list[list[str]] = []
    for result in ordered:
        row: list = [result.record.get(c.field, "") for c in info_columns]
        row += [resolve_value(rule, result.record) for rule in rules]

        reminder = reminder_summary.get(result.asset_id)
        row += [
            next_deadline(config, result),
            reminder.last_sent if reminder else "",
            reminder.count if reminder else 0,
            result.compliance_status,
        ]

        preserved = existing_notes.get(result.asset_id, {})
        row += [preserved.get(c.label, "") or "" for c in notes_columns]

        rows.append(row)

    return headers, ordered, rows


def _write_tracker_sheet(
    ws: Worksheet,
    config: AppConfig,
    results: list[AssetResult],
    reminder_summary: dict[str, ReminderSummary],
    existing_notes: dict[str, dict[str, str]],
) -> None:
    info_columns = config.excel.info_columns
    rules = config.rules
    notes_columns = config.excel.notes_columns

    headers, ordered, rows = build_tracker_table(config, results, reminder_summary, existing_notes)
    _write_header(ws, headers)
    total_cols = len(headers)

    n_info = len(info_columns)
    n_rules = len(rules)
    rule_col_start = n_info + 1
    status_col_idx = n_info + n_rules + 4
    notes_col_start = status_col_idx + 1

    for row_idx, (result, row) in enumerate(zip(ordered, rows), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            if rule_col_start <= col_idx < rule_col_start + n_rules:
                rule = rules[col_idx - rule_col_start]
                violated_rule_ids = {v.rule_id for v in result.violations}
                cell.fill = FAIL_FILL if rule.id in violated_rule_ids else PASS_FILL
            elif col_idx == status_col_idx:
                cell.fill = PASS_FILL if result.compliance_status == COMPLIANT else FAIL_FILL
            elif col_idx >= notes_col_start:
                cell.fill = NOTES_FILL

    _autosize_columns(ws, total_cols)


def _write_audit_log_sheet(ws: Worksheet, results: list[AssetResult]) -> None:
    columns = AUDIT_LOG_COLUMNS
    _write_header(ws, [c.label for c in columns])

    ordered = sorted(results, key=lambda r: r.asset_id)

    row_idx = 2
    for result in ordered:
        for violation in sorted(result.violations, key=lambda v: v.severity != "critical"):
            row_data = {
                "asset_id": result.asset_id,
                "rule_id": violation.rule_id,
                "severity": violation.severity,
                "field": violation.field or "",
                "value": violation.value,
                "message": violation.message,
            }
            for col_idx, column in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(column.field, ""))
            severity_col_idx = next(
                (i for i, c in enumerate(columns, start=1) if c.field == "severity"), None
            )
            if severity_col_idx:
                fill = CRITICAL_FILL if violation.severity == "critical" else WARNING_FILL
                ws.cell(row=row_idx, column=severity_col_idx).fill = fill
            row_idx += 1

    _autosize_columns(ws, len(columns))


def generate_excel_report(
    config: AppConfig,
    results: list[AssetResult],
    reminder_summary: dict[str, ReminderSummary],
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    existing_notes = _read_existing_notes(output_path, config)

    workbook = Workbook()

    ws_tracker = workbook.active
    ws_tracker.title = "Tracker"
    _write_tracker_sheet(ws_tracker, config, results, reminder_summary, existing_notes)

    ws_audit = workbook.create_sheet("Audit Log")
    _write_audit_log_sheet(ws_audit, results)

    workbook.save(output_path)
    return output_path
