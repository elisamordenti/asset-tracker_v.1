from openpyxl import load_workbook

from compliance_tracker.config_schema import (
    AppConfig,
    ColumnConfig,
    ContactConfig,
    EmailConfig,
    ExcelConfig,
    RuleConfig,
    SourceConfig,
)
from compliance_tracker.excel_report import generate_excel_report
from compliance_tracker.reminder_log import ReminderSummary
from compliance_tracker.validator import validate_assets


def build_config(tmp_path, csv_content, rules, info_columns, source_field="asset_id", notes_columns=None):
    csv_path = tmp_path / "assets.csv"
    csv_path.write_text(csv_content, encoding="utf-8")
    return AppConfig(
        domain="Test Domain",
        source=SourceConfig(type="csv", path=str(csv_path), id_field=source_field),
        contact=ContactConfig(name_field="contact_name", email_field="contact_email"),
        rules=rules,
        excel=ExcelConfig(info_columns=info_columns, notes_columns=notes_columns or []),
        email=EmailConfig(subject_template="Subject {" + source_field + "}", body_template="{issues_list}"),
    )


def test_tracker_headers_come_from_config_and_rule_labels(tmp_path):
    csv_content = "asset_id,margin_pct,contact_name,contact_email\nAST-1,2,Dana,dana@example.com\n"
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin Custom Label", params={"min": 5}),
    ]
    info_columns = [ColumnConfig(field="asset_id", label="Totally Custom Label")]
    config = build_config(tmp_path, csv_content, rules, info_columns)

    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    headers = [c.value for c in wb["Tracker"][1]]

    assert headers == [
        "Totally Custom Label", "Margin Custom Label",
        "Next Deadline", "Last Reminder Sent", "Reminder Count", "Status",
    ]


def test_report_has_two_sheets(tmp_path):
    csv_content = "asset_id,contact_name,contact_email\nAST-1,Dana,dana@example.com\n"
    config = build_config(
        tmp_path, csv_content,
        rules=[RuleConfig(id="r", field="asset_id", type="required", severity="warning",
                           message="m", label="R", params={})],
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    assert wb.sheetnames == ["Tracker", "Audit Log"]


def test_tracker_shows_rule_value_regardless_of_pass_fail(tmp_path):
    csv_content = (
        "asset_id,margin_pct,contact_name,contact_email\n"
        "AST-1,10,Dana,dana@example.com\n"
        "AST-2,2,Ben,ben@example.com\n"
    )
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    rows = {row[0]: row[1] for row in wb["Tracker"].iter_rows(min_row=2, values_only=True)}
    assert rows["AST-1"] == "10"
    assert rows["AST-2"] == "2"


def test_tracker_status_column_reflects_overall_compliance(tmp_path):
    csv_content = (
        "asset_id,margin_pct,contact_name,contact_email\n"
        "AST-1,10,Dana,dana@example.com\n"
        "AST-2,2,Ben,ben@example.com\n"
    )
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    rows = {row[0]: row[-1] for row in wb["Tracker"].iter_rows(min_row=2, values_only=True)}
    assert rows["AST-1"] == "COMPLIANT"
    assert rows["AST-2"] == "FLAGGED"


def test_tracker_sorts_flagged_before_compliant(tmp_path):
    csv_content = (
        "asset_id,margin_pct,contact_name,contact_email\n"
        "AST-1,10,Dana,dana@example.com\n"
        "AST-2,1,Ben,ben@example.com\n"
    )
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    rows = list(wb["Tracker"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "AST-2"  # flagged asset listed first
    assert rows[0][-1] == "FLAGGED"
    assert rows[1][-1] == "COMPLIANT"


def test_tracker_includes_reminder_summary_columns(tmp_path):
    csv_content = "asset_id,margin_pct,contact_name,contact_email\nAST-1,2,Dana,dana@example.com\n"
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    reminder_summary = {"AST-1": ReminderSummary(last_sent="2026-08-01", count=3)}
    output_path = generate_excel_report(config, results, reminder_summary, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    headers = [c.value for c in wb["Tracker"][1]]
    row = next(wb["Tracker"].iter_rows(min_row=2, values_only=True))
    row_by_header = dict(zip(headers, row))

    assert row_by_header["Last Reminder Sent"] == "2026-08-01"
    assert row_by_header["Reminder Count"] == 3


def test_tracker_next_deadline_is_earliest_not_expired_date(tmp_path):
    csv_content = (
        "asset_id,cert_expiry,insurance_expiry,contact_name,contact_email\n"
        "AST-1,2027-06-01,2027-01-01,Dana,dana@example.com\n"
    )
    rules = [
        RuleConfig(id="cert_ok", field="cert_expiry", type="not_expired", severity="critical",
                   message="expired", label="Cert", params={"max_age_days": 0}),
        RuleConfig(id="insurance_ok", field="insurance_expiry", type="not_expired", severity="critical",
                   message="expired", label="Insurance", params={"max_age_days": 0}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    headers = [c.value for c in wb["Tracker"][1]]
    row = next(wb["Tracker"].iter_rows(min_row=2, values_only=True))
    row_by_header = dict(zip(headers, row))

    assert row_by_header["Next Deadline"] == "2027-01-01"


def test_audit_log_has_one_row_per_violation(tmp_path):
    csv_content = (
        "asset_id,margin_pct,permit_doc_ref,contact_name,contact_email\n"
        "AST-1,2,,Dana,dana@example.com\n"
    )
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
        RuleConfig(id="has_permit", field="permit_doc_ref", type="required", severity="critical",
                   message="missing permit", label="Permit", params={}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    results = validate_assets(config)
    output_path = generate_excel_report(config, results, {}, tmp_path / "tracker.xlsx")

    wb = load_workbook(output_path)
    data_rows = list(wb["Audit Log"].iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 2
    rule_ids = {row[1] for row in data_rows}
    assert rule_ids == {"margin_min", "has_permit"}


def _notes_config(tmp_path, csv_content, rules):
    return build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
        notes_columns=[ColumnConfig(field="notes", label="Notes")],
    )


def _tracker_row_by_id(path):
    wb = load_workbook(path)
    ws = wb["Tracker"]
    headers = [c.value for c in ws[1]]
    return {row[0]: dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)}


def test_notes_survive_regeneration_for_asset_still_present(tmp_path):
    csv_content = "asset_id,margin_pct,contact_name,contact_email\nAST-1,10,Dana,dana@example.com\n"
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = _notes_config(tmp_path, csv_content, rules)
    output_path = tmp_path / "tracker.xlsx"

    results = validate_assets(config)
    generate_excel_report(config, results, {}, output_path)

    # Simulate a human typing a note directly into the generated file.
    wb = load_workbook(output_path)
    ws = wb["Tracker"]
    headers = [c.value for c in ws[1]]
    notes_col = headers.index("Notes") + 1
    ws.cell(row=2, column=notes_col, value="Called client, awaiting reply")
    wb.save(output_path)

    # Re-run: margin now fails too, so the rule column must refresh --
    # but the note must survive.
    csv_content_v2 = "asset_id,margin_pct,contact_name,contact_email\nAST-1,2,Dana,dana@example.com\n"
    config_v2 = _notes_config(tmp_path, csv_content_v2, rules)
    results_v2 = validate_assets(config_v2)
    generate_excel_report(config_v2, results_v2, {}, output_path)

    row = _tracker_row_by_id(output_path)["AST-1"]
    assert row["Notes"] == "Called client, awaiting reply"
    assert row["Margin"] == "2"
    assert row["Status"] == "FLAGGED"


def test_notes_dropped_for_asset_no_longer_present(tmp_path):
    csv_content = (
        "asset_id,margin_pct,contact_name,contact_email\n"
        "AST-1,10,Dana,dana@example.com\n"
        "AST-2,10,Ben,ben@example.com\n"
    )
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = _notes_config(tmp_path, csv_content, rules)
    output_path = tmp_path / "tracker.xlsx"
    results = validate_assets(config)
    generate_excel_report(config, results, {}, output_path)

    wb = load_workbook(output_path)
    ws = wb["Tracker"]
    headers = [c.value for c in ws[1]]
    notes_col = headers.index("Notes") + 1
    for row in ws.iter_rows(min_row=2):
        row[notes_col - 1].value = "some note"
    wb.save(output_path)

    csv_content_v2 = "asset_id,margin_pct,contact_name,contact_email\nAST-1,10,Dana,dana@example.com\n"
    config_v2 = _notes_config(tmp_path, csv_content_v2, rules)
    results_v2 = validate_assets(config_v2)
    generate_excel_report(config_v2, results_v2, {}, output_path)

    rows = _tracker_row_by_id(output_path)
    assert "AST-2" not in rows
    assert rows["AST-1"]["Notes"] == "some note"


def test_new_asset_starts_with_blank_note(tmp_path):
    csv_content = "asset_id,margin_pct,contact_name,contact_email\nAST-1,10,Dana,dana@example.com\n"
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = _notes_config(tmp_path, csv_content, rules)
    output_path = tmp_path / "tracker.xlsx"
    results = validate_assets(config)
    generate_excel_report(config, results, {}, output_path)

    csv_content_v2 = (
        "asset_id,margin_pct,contact_name,contact_email\n"
        "AST-1,10,Dana,dana@example.com\n"
        "AST-2,10,Ben,ben@example.com\n"
    )
    config_v2 = _notes_config(tmp_path, csv_content_v2, rules)
    results_v2 = validate_assets(config_v2)
    generate_excel_report(config_v2, results_v2, {}, output_path)

    rows = _tracker_row_by_id(output_path)
    assert rows["AST-2"]["Notes"] in (None, "")


def test_no_notes_columns_configured_means_no_read_attempt(tmp_path):
    csv_content = "asset_id,margin_pct,contact_name,contact_email\nAST-1,10,Dana,dana@example.com\n"
    rules = [
        RuleConfig(id="margin_min", field="margin_pct", type="min_value", severity="warning",
                   message="low margin", label="Margin", params={"min": 5}),
    ]
    config = build_config(
        tmp_path, csv_content, rules,
        info_columns=[ColumnConfig(field="asset_id", label="ID")],
    )
    output_path = tmp_path / "tracker.xlsx"
    results = validate_assets(config)
    generate_excel_report(config, results, {}, output_path)
    generate_excel_report(config, results, {}, output_path)  # must not error

    headers = [c.value for c in load_workbook(output_path)["Tracker"][1]]
    assert "Notes" not in headers
