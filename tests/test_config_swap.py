"""Proves the central genericity claim: two structurally different domains
(different field names, different rules, different Excel columns) run
through the exact same validator/excel_report/email_drafter code, driven
only by which config file is passed in."""

from pathlib import Path

from openpyxl import load_workbook

from compliance_tracker.archive import LocalDiskArchive
from compliance_tracker.config_schema import load_config
from compliance_tracker.email_drafter import draft_emails
from compliance_tracker.excel_report import generate_excel_report
from compliance_tracker.reminder_log import append_entries, load_summary
from compliance_tracker.validator import validate_assets

REPO_ROOT = Path(__file__).parent.parent


def run_domain(monkeypatch, tmp_path, config_relpath):
    monkeypatch.chdir(REPO_ROOT)
    config = load_config(config_relpath)
    results = validate_assets(config, archive=LocalDiskArchive())

    log_path = tmp_path / "reminder_log.csv"
    flagged_ids = [r.asset_id for r in results if r.violations]
    append_entries(log_path, flagged_ids)
    reminder_summary = load_summary(log_path)

    excel_path = generate_excel_report(config, results, reminder_summary, tmp_path / "tracker.xlsx")
    drafts = draft_emails(config, results, reminder_summary, tmp_path / "emails")
    return config, results, excel_path, drafts


def test_energy_assets_domain_runs_end_to_end(monkeypatch, tmp_path):
    config, results, excel_path, drafts = run_domain(
        monkeypatch, tmp_path, "config/energy_assets.yaml"
    )
    assert config.domain == "Energy Assets"
    assert len(results) > 0
    assert excel_path.exists()
    assert all(d.file_path.exists() for d in drafts)


def test_portfolio_companies_domain_runs_end_to_end(monkeypatch, tmp_path):
    config, results, excel_path, drafts = run_domain(
        monkeypatch, tmp_path, "config/portfolio_companies.yaml"
    )
    assert config.domain == "Portfolio Companies"
    assert len(results) > 0
    assert excel_path.exists()
    assert all(d.file_path.exists() for d in drafts)


def test_both_domains_produce_domain_specific_excel_columns(monkeypatch, tmp_path):
    """Same generate_excel_report function, different column headers -- proof
    that columns are read from config, not hardcoded per domain."""
    _, _, energy_excel, _ = run_domain(
        monkeypatch, tmp_path / "energy", "config/energy_assets.yaml"
    )
    _, _, portfolio_excel, _ = run_domain(
        monkeypatch, tmp_path / "portfolio", "config/portfolio_companies.yaml"
    )

    energy_headers = [c.value for c in load_workbook(energy_excel)["Tracker"][1]]
    portfolio_headers = [c.value for c in load_workbook(portfolio_excel)["Tracker"][1]]

    assert energy_headers == [
        "Asset ID", "Asset Name", "Location", "Responsible Contact",
        "Grid Cert Expiry", "Grid Cert on File", "Insurance Expiry", "Insurance Cert on File",
        "Permit on File", "Margin %", "Downtime (hrs/yr)", "Maintenance Status",
        "Next Deadline", "Last Reminder Sent", "Reminder Count", "Status", "Notes / Follow-up",
    ]
    assert portfolio_headers == [
        "Company ID", "Company Name", "Stage", "Founder Contact",
        "Financials Audited", "Audit Report on File", "Cap Table Current", "Cap Table Export on File",
        "Last Board Update", "Runway (months)", "Data Room Ref", "Data Room Ref Format", "Funding Stage",
        "Next Deadline", "Last Reminder Sent", "Reminder Count", "Status", "Notes / Follow-up",
    ]
    assert energy_headers != portfolio_headers


def test_both_domains_use_identical_validator_and_excel_functions():
    """Guards against a future regression where someone adds a domain-name
    branch to the engine instead of expressing the difference in config."""
    import inspect

    from compliance_tracker import validator, excel_report

    validate_source = inspect.getsource(validator.validate_assets)
    excel_source = inspect.getsource(excel_report.generate_excel_report)

    for banned in ["energy", "portfolio", "asset_id ==", "company_id =="]:
        assert banned not in validate_source.lower()
        assert banned not in excel_source.lower()
